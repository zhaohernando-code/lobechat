from __future__ import annotations

import json
import os
import re
import shlex
import shutil
import ssl
import subprocess
import hashlib
import time
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse, urlunparse
from urllib.request import Request, urlopen

from fastmcp import FastMCP

BASE_DIR = Path(os.environ.get("LOCAL_OFFICE_MCP_OUTPUT_DIR", "/Users/hernando_zhao/codex/projects/lobechat/output/office-mcp"))
INPUT_CACHE_DIR = Path(os.environ.get("LOCAL_OFFICE_MCP_INPUT_CACHE_DIR", str(BASE_DIR / ".inputs")))
PUBLIC_BASE_URL = os.environ.get("LOCAL_OFFICE_MCP_PUBLIC_BASE_URL", "https://hernando-zhao.cn/chat-files/office")
REMOTE_SYNC_TARGET = os.environ.get("LOCAL_OFFICE_MCP_REMOTE_SYNC_TARGET", "").strip()
AUDIT_LOG = Path(os.environ.get("LOCAL_OFFICE_MCP_AUDIT_LOG", "/Users/hernando_zhao/codex/projects/lobechat/logs/local-office-mcp-audit.log"))
BASE_DIR.mkdir(parents=True, exist_ok=True)
INPUT_CACHE_DIR.mkdir(parents=True, exist_ok=True)
AUDIT_LOG.parent.mkdir(parents=True, exist_ok=True)

mcp = FastMCP("Local Office MCP")

MARKET_SKILL_API = "https://market.lobehub.com/api/v1/skills"
DEFAULT_SKILL_USER_EMAIL = os.environ.get("LOCAL_OFFICE_MCP_DEFAULT_USER_EMAIL", "root@hernando-zhao.cn")
POSTGRES_CONTAINER = os.environ.get("LOCAL_OFFICE_MCP_POSTGRES_CONTAINER", "lobehub-postgresql")
POSTGRES_DB = os.environ.get("LOCAL_OFFICE_MCP_POSTGRES_DB", "lobechat")
POSTGRES_USER = os.environ.get("LOCAL_OFFICE_MCP_POSTGRES_USER", "postgres")
DOCKER_BIN = os.environ.get("LOCAL_OFFICE_MCP_DOCKER_BIN", shutil.which("docker") or "/usr/local/bin/docker")
TESSERACT_BIN = os.environ.get("LOCAL_OFFICE_MCP_TESSERACT_BIN", shutil.which("tesseract") or "/opt/homebrew/bin/tesseract")
try:
    import certifi

    HTTPS_CONTEXT = ssl.create_default_context(cafile=certifi.where())
except Exception:
    HTTPS_CONTEXT = ssl.create_default_context()


def _safe_name(name: str | None, suffix: str) -> str:
    base = (name or "office-output").strip()
    base = re.sub(r"[^A-Za-z0-9._-]+", "-", base).strip(".-") or "office-output"
    if not base.lower().endswith(suffix):
        base += suffix
    return base


def _output(name: str | None, suffix: str) -> Path:
    path = BASE_DIR / _safe_name(name, suffix)
    if path.exists():
        stem = path.stem
        for index in range(2, 1000):
            candidate = BASE_DIR / f"{stem}-{index}{path.suffix}"
            if not candidate.exists():
                return candidate
    return path


def _result(path: Path, summary: str) -> dict[str, Any]:
    return {
        "ok": True,
        "summary": summary,
        "path": str(path),
        "download_url": f"{PUBLIC_BASE_URL.rstrip('/')}/{path.name}",
    }


def _audit(event: str, payload: dict[str, Any]) -> None:
    record = {"event": event, **payload}
    with AUDIT_LOG.open("a", encoding="utf-8") as file:
        file.write(json.dumps(record, ensure_ascii=False) + "\n")


def _sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def _psql(sql: str) -> str:
    command = [
        DOCKER_BIN,
        "exec",
        "-i",
        POSTGRES_CONTAINER,
        "psql",
        "-U",
        POSTGRES_USER,
        "-d",
        POSTGRES_DB,
        "-v",
        "ON_ERROR_STOP=1",
        "-t",
        "-A",
        "-F",
        "\t",
    ]
    completed = subprocess.run(command, input=sql, capture_output=True, text=True, timeout=30)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or completed.stdout.strip() or "psql failed")
    return completed.stdout.strip()


def _market_skill(identifier: str) -> dict[str, Any]:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", identifier.strip()).strip(".-")
    if not cleaned:
        raise ValueError("identifier is required")
    request = Request(f"{MARKET_SKILL_API}/{cleaned}", headers={"User-Agent": "local-office-mcp/1.1"})
    with urlopen(request, timeout=30, context=HTTPS_CONTEXT) as response:
        return json.loads(response.read().decode("utf-8"))


def _download_market_skill_zip(identifier: str) -> tuple[str, int]:
    request = Request(f"{MARKET_SKILL_API}/{identifier}/download", headers={"User-Agent": "local-office-mcp/1.1"})
    with urlopen(request, timeout=30, context=HTTPS_CONTEXT) as response:
        data = response.read()
    return hashlib.sha256(data).hexdigest(), len(data)


def _agent_plugin_append_sql(user_ids: list[str], identifier: str) -> str:
    quoted_users = ", ".join(_sql_literal(user_id) for user_id in user_ids)
    quoted_identifier = _sql_literal(identifier)
    return f"""
WITH target_agents AS (
  SELECT id, plugins
  FROM agents
  WHERE user_id IN ({quoted_users})
    AND provider = 'deepseek'
),
normalized AS (
  SELECT
    id,
    (
      SELECT jsonb_agg(value)
      FROM (
        SELECT DISTINCT value
        FROM jsonb_array_elements_text(COALESCE(target_agents.plugins, '[]'::jsonb) || jsonb_build_array({quoted_identifier})) AS t(value)
        WHERE value <> 'lobe-cloud-sandbox'
      ) AS unique_values
    ) AS next_plugins
  FROM target_agents
)
UPDATE agents
SET plugins = normalized.next_plugins,
    updated_at = now()
FROM normalized
WHERE agents.id = normalized.id;

UPDATE user_settings
SET default_agent = jsonb_set(
      COALESCE(default_agent, '{{}}'::jsonb),
      '{{plugins}}',
      (
        SELECT jsonb_agg(value)
        FROM (
          SELECT DISTINCT value
          FROM jsonb_array_elements_text(COALESCE(default_agent->'plugins', '[]'::jsonb) || jsonb_build_array({quoted_identifier})) AS t(value)
          WHERE value <> 'lobe-cloud-sandbox'
        ) AS unique_values
      ),
      true
    )
WHERE id IN ({quoted_users});
"""


def _sync_remote(path: Path) -> dict[str, Any] | None:
    if not REMOTE_SYNC_TARGET:
        return None

    if ":" not in REMOTE_SYNC_TARGET:
        result = {"ok": False, "error": "LOCAL_OFFICE_MCP_REMOTE_SYNC_TARGET must be user@host:/path or host:/path."}
        _audit("remote_sync.error", {"path": str(path), **result})
        return result

    host, remote_dir = REMOTE_SYNC_TARGET.split(":", 1)
    if not host or not remote_dir.startswith("/"):
        result = {"ok": False, "error": "LOCAL_OFFICE_MCP_REMOTE_SYNC_TARGET must use an absolute remote path."}
        _audit("remote_sync.error", {"path": str(path), **result})
        return result

    remote_file = f"{host}:{remote_dir.rstrip('/')}/{path.name}"
    try:
        subprocess.run(
            ["ssh", host, f"mkdir -p {shlex.quote(remote_dir)}"],
            check=True,
            capture_output=True,
            text=True,
            timeout=20,
        )
        subprocess.run(
            ["scp", str(path), remote_file],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
    except subprocess.SubprocessError as error:
        result = {"ok": False, "error": str(error)}
        _audit("remote_sync.error", {"path": str(path), "target": remote_file, **result})
        return result

    result = {"ok": True, "target": remote_file}
    _audit("remote_sync.done", {"path": str(path), **result})
    return result


def _normalize_lines(value: str | list[str] | None) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    return [line.strip() for line in str(value).splitlines() if line.strip()]


def _safe_basename(value: str, fallback: str = "input") -> str:
    parsed_name = unquote(Path(urlparse(value).path).name if "://" in value else Path(value).name)
    name = re.sub(r"[^A-Za-z0-9._-]+", "-", parsed_name).strip(".-")
    return name or fallback


def _local_rustfs_url(parsed_url) -> str:
    path = parsed_url.path
    if path.startswith("/chat-s3/"):
        path = "/" + path[len("/chat-s3/"):]
    return urlunparse(("http", "127.0.0.1:9000", path, "", parsed_url.query, parsed_url.fragment))


def _download_url(value: str) -> Path:
    parsed_url = urlparse(value)
    if parsed_url.hostname in {"hernando-zhao.cn", "www.hernando-zhao.cn"} and parsed_url.path.startswith("/chat-files/office/"):
        candidate = BASE_DIR / _safe_basename(parsed_url.path)
        if candidate.exists():
            return candidate

    fetch_url = value
    if (
        parsed_url.hostname in {"hernando-zhao.cn", "www.hernando-zhao.cn"}
        and parsed_url.path.startswith("/chat-s3/")
    ):
        fetch_url = _local_rustfs_url(parsed_url)

    suffix = Path(urlparse(fetch_url).path).suffix or Path(_safe_basename(value)).suffix or ".bin"
    digest = hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]
    cache_path = INPUT_CACHE_DIR / f"{digest}{suffix}"
    if cache_path.exists():
        return cache_path

    request = Request(fetch_url, headers={"User-Agent": "local-office-mcp/1.0"})
    with urlopen(request, timeout=30) as response:
        cache_path.write_bytes(response.read())
    return cache_path


def _resolve_input_path(location: str) -> Path:
    value = str(location or "").strip()
    if not value:
        raise ValueError("path or URL is required")

    if value.startswith("/chat-files/office/"):
        return BASE_DIR / _safe_basename(value)

    if value.startswith("/chat-s3/"):
        return _download_url(f"http://127.0.0.1:9000/{value[len('/chat-s3/'):].lstrip('/')}")

    if value.startswith("http://") or value.startswith("https://"):
        return _download_url(value)

    path = Path(value)
    if not path.is_absolute():
        path = BASE_DIR / value
    return path


def _read_text(path: Path, encoding: str = "utf-8", max_chars: int = 20000) -> dict[str, Any]:
    data = path.read_bytes()
    text = data.decode(encoding or "utf-8", errors="replace")
    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]
    return {
        "ok": True,
        "path": str(path),
        "encoding": encoding or "utf-8",
        "truncated": truncated,
        "text": text,
    }


@mcp.tool()
def create_docx(title: str, content: str | list[str], output_filename: str | None = None) -> dict[str, Any]:
    """Create a Word .docx document from a title and paragraphs."""
    from docx import Document

    _audit("create_docx.start", {"title": title, "output_filename": output_filename})
    path = _output(output_filename or title, ".docx")
    document = Document()
    document.add_heading(title, 0)
    for line in _normalize_lines(content):
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        else:
            document.add_paragraph(line)
    document.save(path)
    result = _result(path, f"Created Word document with {len(_normalize_lines(content))} content blocks.")
    remote_sync = _sync_remote(path)
    if remote_sync is not None:
        result["remote_sync"] = remote_sync
    _audit("create_docx.done", result)
    return result


@mcp.tool()
def read_docx(path: str) -> dict[str, Any]:
    """Read text from a Word .docx document by local path, generated filename, or /chat-s3 URL."""
    from docx import Document

    _audit("read_docx.start", {"path": path})
    resolved_path = _resolve_input_path(path)
    document = Document(resolved_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    result = {"ok": True, "path": str(resolved_path), "paragraphs": paragraphs, "text": "\n".join(paragraphs)}
    _audit("read_docx.done", {"path": str(resolved_path), "paragraph_count": len(paragraphs)})
    return result


@mcp.tool()
def create_pptx(title: str, slides: list[dict[str, Any]], output_filename: str | None = None) -> dict[str, Any]:
    """Create a PowerPoint .pptx deck. Each slide accepts title and bullets fields."""
    from pptx import Presentation

    _audit("create_pptx.start", {"title": title, "output_filename": output_filename, "slides": len(slides)})
    path = _output(output_filename or title, ".pptx")
    presentation = Presentation()
    for slide_data in slides:
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = str(slide_data.get("title", "Untitled"))
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = slide_data.get("bullets", [])
        if isinstance(bullets, str):
            bullets = _normalize_lines(bullets)
        for index, bullet in enumerate(bullets or []):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.level = 0
    presentation.save(path)
    result = _result(path, f"Created PowerPoint deck with {len(slides)} slides.")
    remote_sync = _sync_remote(path)
    if remote_sync is not None:
        result["remote_sync"] = remote_sync
    _audit("create_pptx.done", result)
    return result


@mcp.tool()
def read_pptx(path: str) -> dict[str, Any]:
    """Read titles and visible text from a PowerPoint .pptx deck by local path, generated filename, or /chat-s3 URL."""
    from pptx import Presentation

    _audit("read_pptx.start", {"path": path})
    resolved_path = _resolve_input_path(path)
    presentation = Presentation(resolved_path)
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        slides.append({"slide": index, "text": "\n".join(texts)})
    result = {"ok": True, "path": str(resolved_path), "slides": slides}
    _audit("read_pptx.done", {"path": str(resolved_path), "slide_count": len(slides)})
    return result


@mcp.tool()
def create_xlsx(workbook_name: str, rows: list[list[Any]], output_filename: str | None = None) -> dict[str, Any]:
    """Create an Excel .xlsx workbook from row data."""
    from openpyxl import Workbook

    _audit("create_xlsx.start", {"workbook_name": workbook_name, "output_filename": output_filename, "rows": len(rows)})
    path = _output(output_filename or workbook_name, ".xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = workbook_name[:31] or "Sheet1"
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    result = _result(path, f"Created Excel workbook with {len(rows)} rows.")
    remote_sync = _sync_remote(path)
    if remote_sync is not None:
        result["remote_sync"] = remote_sync
    _audit("create_xlsx.done", result)
    return result


@mcp.tool()
def read_xlsx(path: str, max_rows: int = 100) -> dict[str, Any]:
    """Read values from the first sheet of an Excel .xlsx workbook by local path, generated filename, or /chat-s3 URL."""
    from openpyxl import load_workbook

    _audit("read_xlsx.start", {"path": path, "max_rows": max_rows})
    resolved_path = _resolve_input_path(path)
    workbook = load_workbook(resolved_path, data_only=True, read_only=True)
    sheet = workbook.active
    rows: list[list[Any]] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index > max_rows:
            break
        rows.append(list(row))
    result = {"ok": True, "path": str(resolved_path), "sheet": sheet.title, "rows": rows}
    _audit("read_xlsx.done", {"path": str(resolved_path), "row_count": len(rows)})
    return result


@mcp.tool()
def read_text_file(path: str, encoding: str = "utf-8", max_chars: int = 20000) -> dict[str, Any]:
    """Read text, code, Markdown, CSV, JSON, or other plaintext files by local path, generated filename, or /chat-s3 URL."""
    _audit("read_text_file.start", {"path": path, "encoding": encoding, "max_chars": max_chars})
    resolved_path = _resolve_input_path(path)
    result = _read_text(resolved_path, encoding=encoding, max_chars=max_chars)
    _audit("read_text_file.done", {"path": str(resolved_path), "chars": len(result["text"]), "truncated": result["truncated"]})
    return result


@mcp.tool()
def read_file(path: str, max_chars: int = 20000, max_rows: int = 100) -> dict[str, Any]:
    """Read a supported text, Word, PowerPoint, or Excel file by local path, generated filename, or /chat-s3 URL."""
    _audit("read_file.start", {"path": path, "max_chars": max_chars, "max_rows": max_rows})
    resolved_path = _resolve_input_path(path)
    suffix = resolved_path.suffix.lower()
    if suffix == ".docx":
        return read_docx(str(resolved_path))
    if suffix == ".pptx":
        return read_pptx(str(resolved_path))
    if suffix == ".xlsx":
        return read_xlsx(str(resolved_path), max_rows=max_rows)
    result = _read_text(resolved_path, max_chars=max_chars)
    _audit("read_file.done", {"path": str(resolved_path), "kind": "text", "chars": len(result["text"])})
    return result


@mcp.tool()
def ocr_image(path: str, lang: str = "eng+chi_sim") -> dict[str, Any]:
    """Run OCR on an image file using local Tesseract."""
    _audit("ocr_image.start", {"path": path, "lang": lang})
    if not Path(TESSERACT_BIN).exists():
        result = {
            "ok": False,
            "error": "Tesseract binary is not installed on the host.",
            "hint": "Install tesseract and language packs, then retry.",
        }
        _audit("ocr_image.error", result)
        return result
    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = TESSERACT_BIN
    resolved_path = _resolve_input_path(path)
    text = pytesseract.image_to_string(Image.open(resolved_path), lang=lang)
    result = {"ok": True, "path": str(resolved_path), "text": text}
    _audit("ocr_image.done", {"path": str(resolved_path), "chars": len(text)})
    return result


@mcp.tool()
def list_outputs() -> dict[str, Any]:
    """List files generated by the local Office MCP server."""
    _audit("list_outputs.start", {})
    files = []
    for path in sorted(BASE_DIR.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if path.is_file():
            files.append({"name": path.name, "path": str(path), "download_url": f"{PUBLIC_BASE_URL.rstrip('/')}/{path.name}"})
    result = {"ok": True, "files": files}
    _audit("list_outputs.done", {"count": len(files)})
    return result


@mcp.tool()
def get_market_skill(identifier: str) -> dict[str, Any]:
    """Fetch one LobeHub market skill by identifier without using Cloud Sandbox authorization."""
    _audit("get_market_skill.start", {"identifier": identifier})
    skill = _market_skill(identifier)
    result = {
        "ok": True,
        "identifier": skill.get("identifier"),
        "name": skill.get("name"),
        "description": skill.get("description"),
        "version": skill.get("version"),
        "install_count": skill.get("installCount"),
        "manifest": skill.get("manifest") or {},
    }
    _audit("get_market_skill.done", {"identifier": result["identifier"], "name": result["name"]})
    return result


@mcp.tool()
def list_installed_skills(user_email: str = DEFAULT_SKILL_USER_EMAIL) -> dict[str, Any]:
    """List skills locally installed in LobeChat for a user."""
    _audit("list_installed_skills.start", {"user_email": user_email})
    output = _psql(
        f"""
SELECT s.identifier || E'\t' || s.name || E'\t' || s.source || E'\t' || COALESCE(s.updated_at::text, '')
FROM agent_skills s
JOIN users u ON u.id = s.user_id
WHERE u.email = {_sql_literal(user_email)}
ORDER BY s.updated_at DESC;
"""
    )
    skills = []
    for line in output.splitlines():
        if not line.strip():
            continue
        identifier, name, source, updated_at = (line.split("\t") + ["", "", "", ""])[:4]
        skills.append({"identifier": identifier, "name": name, "source": source, "updated_at": updated_at})
    result = {"ok": True, "user_email": user_email, "skills": skills}
    _audit("list_installed_skills.done", {"user_email": user_email, "count": len(skills)})
    return result


@mcp.tool()
def import_market_skill(
    identifier: str,
    user_email: str = DEFAULT_SKILL_USER_EMAIL,
    install_for_all_users: bool = False,
    attach_to_deepseek: bool = True,
) -> dict[str, Any]:
    """Import a LobeHub market skill into local LobeChat tables and optionally attach it to DeepSeek agents."""
    _audit(
        "import_market_skill.start",
        {
            "identifier": identifier,
            "user_email": user_email,
            "install_for_all_users": install_for_all_users,
            "attach_to_deepseek": attach_to_deepseek,
        },
    )
    skill = _market_skill(identifier)
    skill_identifier = str(skill.get("identifier") or identifier)
    skill_name = str(skill.get("name") or skill_identifier)
    manifest = skill.get("manifest") or {}
    manifest.setdefault("name", skill_name)
    manifest.setdefault("description", skill.get("description") or "")
    manifest.setdefault("sourceUrl", f"{MARKET_SKILL_API}/{skill_identifier}/download")
    content = skill.get("content") or ""

    users_sql = "SELECT id, email FROM users"
    if not install_for_all_users:
        users_sql += f" WHERE email = {_sql_literal(user_email)}"
    users_output = _psql(users_sql + " ORDER BY created_at;")
    users = []
    for line in users_output.splitlines():
        if not line.strip():
            continue
        user_id, email = (line.split("\t") + [""])[:2]
        users.append({"id": user_id, "email": email})
    if not users:
        result = {"ok": False, "error": f"No LobeChat user matched {user_email!r}."}
        _audit("import_market_skill.error", result)
        return result

    zip_hash = None
    zip_size = None
    try:
        zip_hash, zip_size = _download_market_skill_zip(skill_identifier)
    except Exception as error:  # zip is useful for audit, but content+manifest are enough to run a skill.
        _audit("import_market_skill.zip_warning", {"identifier": skill_identifier, "error": str(error)})

    manifest_json = json.dumps(manifest, ensure_ascii=False)
    resources_json = json.dumps(skill.get("resources") or {}, ensure_ascii=False)
    installed = []
    for user in users:
        row_id = "skl_local_" + hashlib.sha256(f"{user['id']}:{skill_identifier}".encode("utf-8")).hexdigest()[:16]
        _psql(
            f"""
INSERT INTO agent_skills (
  id, name, description, identifier, source, manifest, content, resources, zip_file_hash, user_id, accessed_at, created_at, updated_at
)
VALUES (
  {_sql_literal(row_id)},
  {_sql_literal(skill_name)},
  {_sql_literal(skill.get("description") or "")},
  {_sql_literal(skill_identifier)},
  'market',
  {_sql_literal(manifest_json)}::jsonb,
  {_sql_literal(content)},
  {_sql_literal(resources_json)}::jsonb,
  NULL,
  {_sql_literal(user["id"])},
  now(),
  now(),
  now()
)
ON CONFLICT (user_id, name)
DO UPDATE SET
  description = EXCLUDED.description,
  identifier = EXCLUDED.identifier,
  source = EXCLUDED.source,
  manifest = EXCLUDED.manifest,
  content = EXCLUDED.content,
  resources = EXCLUDED.resources,
  accessed_at = now(),
  updated_at = now();
"""
        )
        installed.append({"user_id": user["id"], "email": user["email"]})

    if attach_to_deepseek:
        _psql(_agent_plugin_append_sql([user["id"] for user in users], skill_identifier))

    result = {
        "ok": True,
        "status": "installed_or_updated",
        "identifier": skill_identifier,
        "name": skill_name,
        "users": installed,
        "attached_to_deepseek": attach_to_deepseek,
        "downloaded_zip_sha256": zip_hash,
        "downloaded_zip_size": zip_size,
        "note": "Installed locally; no LobeHub Cloud Sandbox authorization was used.",
    }
    _audit("import_market_skill.done", result)
    return result


if __name__ == "__main__":
    port = int(os.environ.get("LOCAL_OFFICE_MCP_PORT", "18081"))
    mcp.run(transport="streamable-http", host="127.0.0.1", port=port, path="/mcp", stateless_http=True)
